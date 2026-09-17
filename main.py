#!/usr/bin/env python3
"""
Irish Rail Spend Tracker
========================
Scans your Gmail for Irish Rail booking confirmation emails, filters by
passenger name, excludes cancelled bookings, and produces a week-by-week
Excel spreadsheet showing your travel and spend.

Usage:
    python main.py

See README.md for setup instructions.
"""

import os
import re
import json
import pickle
import argparse
from datetime import datetime, timedelta, date
from collections import defaultdict

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
import base64
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── Gmail API scope (read-only) ───────────────────────────────────────────────
SCOPES = ["https://www.googleapis.com/auth/gmail.readonly"]

IRISH_RAIL_SENDER = "iebookinginfo@irishrail.ie"


# ── Auth ──────────────────────────────────────────────────────────────────────

def get_gmail_service():
    """Authenticate with Gmail API and return a service object."""
    creds = None
    token_path = "token.pickle"
    creds_path = "credentials.json"

    if not os.path.exists(creds_path):
        raise FileNotFoundError(
            "credentials.json not found.\n"
            "Please follow the setup instructions in README.md to download "
            "your Google OAuth credentials."
        )

    if os.path.exists(token_path):
        with open(token_path, "rb") as f:
            creds = pickle.load(f)

    if not creds or not creds.valid:
        if creds and creds.expired and creds.refresh_token:
            creds.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file(creds_path, SCOPES)
            creds = flow.run_local_server(port=0)
        with open(token_path, "wb") as f:
            pickle.dump(creds, f)

    return build("gmail", "v1", credentials=creds)


# ── Gmail helpers ─────────────────────────────────────────────────────────────

def search_messages(service, query, max_results=500):
    """Return list of message dicts matching the query."""
    messages = []
    page_token = None

    while True:
        kwargs = {"userId": "me", "q": query, "maxResults": min(500, max_results)}
        if page_token:
            kwargs["pageToken"] = page_token

        result = service.users().messages().list(**kwargs).execute()
        batch = result.get("messages", [])
        messages.extend(batch)

        page_token = result.get("nextPageToken")
        if not page_token or len(messages) >= max_results:
            break

    return messages[:max_results]


def get_message_body(service, msg_id):
    """Fetch the plain-text body of a Gmail message."""
    msg = service.users().messages().get(
        userId="me", id=msg_id, format="full"
    ).execute()

    def extract_parts(payload):
        """Recursively extract text/plain parts."""
        mime = payload.get("mimeType", "")
        if mime == "text/plain":
            data = payload.get("body", {}).get("data", "")
            if data:
                return base64.urlsafe_b64decode(data).decode("utf-8", errors="replace")
        for part in payload.get("parts", []):
            result = extract_parts(part)
            if result:
                return result
        return ""

    body = extract_parts(msg.get("payload", {}))

    # Also grab internal date for fallback
    internal_date = int(msg.get("internalDate", 0)) / 1000
    email_date = datetime.fromtimestamp(internal_date) if internal_date else datetime.now()

    return body, email_date


# ── Parsing ───────────────────────────────────────────────────────────────────

def extract_booking_ref(body):
    """Extract booking reference number from email body."""
    m = re.search(
        r"[Bb]ooking\s+(?:reference\s+)?(?:number|ref)[^\d]*(\d{6,12})",
        body, re.IGNORECASE
    )
    return m.group(1) if m else ""


def norm_name(s):
    """Normalise a name for comparison (lowercase, remove accents)."""
    import unicodedata
    s = unicodedata.normalize("NFD", s.lower())
    return "".join(c for c in s if unicodedata.category(c) != "Mn")


def is_target_passenger(name, target_norm):
    """Check if a passenger name matches the target."""
    return target_norm in norm_name(name)


def parse_booking_email(body, email_date, msg_id, target_norm):
    """
    Parse a booking confirmation email.
    Returns a booking dict if the target passenger is listed, else None.
    """
    if not body:
        return None

    booking_ref = extract_booking_ref(body)

    # ── Passengers section ────────────────────────────────────────────────────
    pax_match = re.search(
        r"Passengers([\s\S]*?)(?=Payment information|Legend|$)",
        body, re.IGNORECASE
    )
    if not pax_match:
        return None

    pax_text = pax_match.group(0)
    row_re = re.compile(
        r"\|\s*([A-Z][a-zA-Z\u00C0-\u017E\-' ]+)\s*\|\s*"
        r"(Adult|Child|Senior|Young Adult|Student)[^|]*\|",
        re.IGNORECASE
    )

    passengers = []
    for m in row_re.finditer(pax_text):
        name = m.group(1).strip()
        ptype = m.group(2).strip()
        if name and not re.match(r"^(Full Name|Passenger)", name, re.IGNORECASE):
            passengers.append({"name": name, "type": ptype})

    if not passengers:
        return None

    target_idx = next(
        (i for i, p in enumerate(passengers) if is_target_passenger(p["name"], target_norm)),
        None
    )
    if target_idx is None:
        return None

    # ── Journey details ───────────────────────────────────────────────────────
    out_match = re.search(
        r"Outward\s+([A-Za-z ]+?)\s+to\s+([A-Za-z ]+?)\s+"
        r"((?:Mon|Tue|Wed|Thu|Fri|Sat|Sun)\s+\d+\s+\w+,?\s*\d{4})",
        body, re.IGNORECASE
    )
    travel_date_str = out_match.group(3).strip() if out_match else ""
    origin = out_match.group(1).strip() if out_match else ""
    destination = out_match.group(2).strip() if out_match else ""

    # Parse travel date
    travel_date = None
    if travel_date_str:
        cleaned = travel_date_str.replace(",", "").strip()
        for fmt in ("%a %d %b %Y", "%A %d %B %Y", "%a %d %B %Y"):
            try:
                travel_date = datetime.strptime(cleaned, fmt)
                break
            except ValueError:
                # try stripping weekday
                parts = cleaned.split(" ", 1)
                if len(parts) > 1:
                    try:
                        travel_date = datetime.strptime(parts[1], fmt.split(" ", 1)[1])
                        break
                    except ValueError:
                        continue

    if not travel_date:
        travel_date = email_date

    # ── Amount — prorated by passenger count ──────────────────────────────────
    total_match = re.search(r"Total paid\s*\|\s*€([\d.,]+)", body, re.IGNORECASE)
    passenger_count = len(passengers)
    amount_raw = 0.0
    amount_display = "—"

    if total_match:
        total = float(total_match.group(1).replace(",", ""))
        amount_raw = total / passenger_count
        amount_display = f"€{amount_raw:.2f}"
        if passenger_count > 1:
            amount_display += f" (1/{passenger_count} of €{total:.2f})"

    other_pax = [
        f"{p['name']} ({p['type']})"
        for i, p in enumerate(passengers) if i != target_idx
    ]

    return {
        "booking_ref": booking_ref,
        "travel_date": travel_date,
        "travel_date_str": travel_date_str,
        "origin": origin,
        "destination": destination,
        "passenger_count": passenger_count,
        "amount_display": amount_display,
        "amount_raw": amount_raw,
        "other_passengers": "; ".join(other_pax),
        "email_date": email_date,
        "msg_id": msg_id,
    }


# ── Week helpers ──────────────────────────────────────────────────────────────

def week_start(dt):
    """Return the Monday of the week containing dt."""
    d = dt.date() if isinstance(dt, datetime) else dt
    return d - timedelta(days=d.weekday())


def week_label(monday):
    """Return a human-readable week label."""
    sunday = monday + timedelta(days=6)
    return f"{monday.strftime('%d %b %Y')} – {sunday.strftime('%d %b %Y')}"


# ── Excel output ──────────────────────────────────────────────────────────────

HEADER_FILL   = PatternFill("solid", fgColor="0052CC")
TRAVEL_FILL   = PatternFill("solid", fgColor="E8F4FD")
NO_TRAVEL_FILL = PatternFill("solid", fgColor="F8F9FA")
YES_FONT      = Font(bold=True, color="065F46")
HEADER_FONT   = Font(bold=True, color="FFFFFF")
THIN_BORDER   = Border(
    left=Side(style="thin", color="DEE2E6"),
    right=Side(style="thin", color="DEE2E6"),
    top=Side(style="thin", color="DEE2E6"),
    bottom=Side(style="thin", color="DEE2E6"),
)


def write_excel(weeks, max_journeys, passenger_name, output_path):
    """Write the week-by-week data to an Excel file."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Weekly Travel"

    # ── Build headers ─────────────────────────────────────────────────────────
    headers = ["Week", "Travelled?", "Week Spend (€)"]
    for i in range(1, max_journeys + 1):
        headers += [f"J{i} Date", f"J{i} Route", f"J{i} My Amount", f"J{i} Booking Ref"]

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        cell.border = THIN_BORDER

    ws.row_dimensions[1].height = 30

    # ── Write rows ────────────────────────────────────────────────────────────
    for row_idx, week in enumerate(weeks, 2):
        monday = week["monday"]
        bookings = week["bookings"]
        travelled = len(bookings) > 0
        week_spend = sum(b["amount_raw"] for b in bookings)

        row_fill = TRAVEL_FILL if travelled else NO_TRAVEL_FILL

        def wcell(col, value, **kwargs):
            cell = ws.cell(row=row_idx, column=col, value=value)
            cell.fill = row_fill
            cell.border = THIN_BORDER
            cell.alignment = Alignment(horizontal=kwargs.get("align", "left"), vertical="center")
            if kwargs.get("bold"):
                cell.font = Font(bold=True)
            if kwargs.get("color"):
                cell.font = Font(bold=kwargs.get("bold", False), color=kwargs["color"])
            return cell

        col = 1
        wcell(col, week_label(monday), bold=travelled); col += 1
        cell = wcell(col, "Yes" if travelled else "—", bold=travelled,
                     color="065F46" if travelled else "6C757D", align="center"); col += 1
        wcell(col, f"€{week_spend:.2f}" if travelled else "—",
              bold=travelled, align="right"); col += 1

        for i in range(max_journeys):
            if i < len(bookings):
                b = bookings[i]
                day_label = b["travel_date"].strftime("%a %d %b") if b["travel_date"] else b["travel_date_str"]
                route = f"{b['origin']} → {b['destination']}" if b["origin"] and b["destination"] else ""
                wcell(col, day_label); col += 1
                wcell(col, route); col += 1
                wcell(col, b["amount_display"], align="right"); col += 1
                ref_cell = wcell(col, b["booking_ref"])
                ref_cell.font = Font(color="6C757D", size=9)
                col += 1
            else:
                for _ in range(4):
                    wcell(col, ""); col += 1

    # ── Summary sheet ─────────────────────────────────────────────────────────
    ws2 = wb.create_sheet("Summary")
    total_spend = sum(b["amount_raw"] for w in weeks for b in w["bookings"])
    weeks_travelled = sum(1 for w in weeks if w["bookings"])
    total_journeys = sum(len(w["bookings"]) for w in weeks)

    summary_data = [
        ("Passenger", passenger_name),
        ("Total weeks analysed", len(weeks)),
        ("Weeks with travel", weeks_travelled),
        ("Weeks without travel", len(weeks) - weeks_travelled),
        ("Total journeys", total_journeys),
        ("Total spend (your share)", f"€{total_spend:.2f}"),
        ("Average spend per week travelled", f"€{total_spend/weeks_travelled:.2f}" if weeks_travelled else "—"),
        ("Average spend per journey", f"€{total_spend/total_journeys:.2f}" if total_journeys else "—"),
    ]

    for r, (label, value) in enumerate(summary_data, 1):
        lc = ws2.cell(row=r, column=1, value=label)
        lc.font = Font(bold=True)
        lc.fill = PatternFill("solid", fgColor="E8F4FD")
        lc.border = THIN_BORDER
        vc = ws2.cell(row=r, column=2, value=value)
        vc.border = THIN_BORDER

    ws2.column_dimensions["A"].width = 35
    ws2.column_dimensions["B"].width = 25

    # ── Column widths (main sheet) ────────────────────────────────────────────
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 16
    for i in range(max_journeys):
        base = 4 + i * 4
        ws.column_dimensions[get_column_letter(base)].width = 16     # date
        ws.column_dimensions[get_column_letter(base+1)].width = 32   # route
        ws.column_dimensions[get_column_letter(base+2)].width = 24   # amount
        ws.column_dimensions[get_column_letter(base+3)].width = 14   # ref

    ws.freeze_panes = "A2"

    wb.save(output_path)


# ── Main ──────────────────────────────────────────────────────────────────────

def main():
    parser = argparse.ArgumentParser(
        description="Irish Rail Spend Tracker — analyse your Irish Rail bookings from Gmail"
    )
    parser.add_argument(
        "--name", "-n",
        help="Passenger name to filter for (e.g. 'Sebastien Marchand')",
        default=None
    )
    parser.add_argument(
        "--months", "-m",
        help="Number of months to look back (default: 13)",
        type=int, default=13
    )
    parser.add_argument(
        "--output", "-o",
        help="Output Excel file path (default: irish_rail_travel.xlsx)",
        default="irish_rail_travel.xlsx"
    )
    args = parser.parse_args()

    # ── Passenger name ────────────────────────────────────────────────────────
    passenger_name = args.name
    if not passenger_name:
        passenger_name = input("Enter the passenger name to search for: ").strip()
    if not passenger_name:
        print("Error: passenger name is required.")
        return

    target_norm = norm_name(passenger_name)

    # ── Date range ────────────────────────────────────────────────────────────
    end_date = datetime.now()
    start_date = end_date - timedelta(days=args.months * 30)
    after_str = start_date.strftime("%Y/%m/%d")
    print(f"\n🔍 Searching emails from {start_date.strftime('%d %b %Y')} to today")
    print(f"👤 Passenger: {passenger_name}\n")

    # ── Connect to Gmail ──────────────────────────────────────────────────────
    print("🔐 Authenticating with Gmail...")
    service = get_gmail_service()
    print("✅ Gmail connected.\n")

    # ── Step 1: Find cancelled booking refs ───────────────────────────────────
    print("🔍 Step 1/2 — Finding cancelled bookings...")
    cancelled_refs = set()
    cancel_query = f"from:{IRISH_RAIL_SENDER} subject:cancelled after:{after_str}"
    cancel_msgs = search_messages(service, cancel_query)
    print(f"   Found {len(cancel_msgs)} possible cancellation emails...")

    for i, msg in enumerate(cancel_msgs):
        try:
            body, _ = get_message_body(service, msg["id"])
            if not re.search(r"your booking has been cancelled", body, re.IGNORECASE):
                continue
            ref = extract_booking_ref(body)
            if ref and ref not in cancelled_refs:
                cancelled_refs.add(ref)
                print(f"   🚫 Cancelled: {ref}")
        except Exception as e:
            pass  # skip unreadable messages

    print(f"\n   {len(cancelled_refs)} cancelled booking(s) will be excluded.\n")

    # ── Step 2: Find and parse booking confirmations ──────────────────────────
    print("📬 Step 2/2 — Fetching booking confirmation emails...")
    booking_query = (
        f"from:{IRISH_RAIL_SENDER} "
        f"subject:\"Thank you for booking\" "
        f"after:{after_str}"
    )
    booking_msgs = search_messages(service, booking_query)
    print(f"   Found {len(booking_msgs)} booking emails. Checking each...\n")

    bookings = []
    matched = 0
    excluded_cancelled = 0
    no_passenger = 0

    for i, msg in enumerate(booking_msgs):
        print(f"   [{i+1}/{len(booking_msgs)}] Reading...", end="\r")
        try:
            body, email_date = get_message_body(service, msg["id"])
            booking = parse_booking_email(body, email_date, msg["id"], target_norm)

            if not booking:
                no_passenger += 1
                continue

            if booking["booking_ref"] and booking["booking_ref"] in cancelled_refs:
                excluded_cancelled += 1
                print(f"   🚫 [{booking['booking_ref']}] Cancelled — excluded")
                continue

            bookings.append(booking)
            matched += 1
            route = f"{booking['origin']} → {booking['destination']}" if booking["origin"] else ""
            print(f"   ✅ [{booking['booking_ref']}] {booking['travel_date_str']} {route} {booking['amount_display']}")

        except Exception as e:
            print(f"   ⚠️  Error reading message {msg['id']}: {e}")

    print(f"\n{'─'*60}")
    print(f"✅ Done! {matched} journeys found for {passenger_name}")
    print(f"🚫 {excluded_cancelled} cancelled booking(s) excluded")
    print(f"⏭  {no_passenger} emails without {passenger_name} as passenger")

    if not bookings:
        print("\nNo bookings found. Check the passenger name and try again.")
        return

    # ── Build week grid ───────────────────────────────────────────────────────
    # Sort by travel date
    bookings.sort(key=lambda b: b["travel_date"])

    # Group by week
    by_week = defaultdict(list)
    for b in bookings:
        mon = week_start(b["travel_date"])
        by_week[mon].append(b)

    # Generate every week from start_date to today
    weeks = []
    current = week_start(start_date)
    today_week = week_start(datetime.now())
    while current <= today_week:
        weeks.append({
            "monday": current,
            "bookings": by_week.get(current, [])
        })
        current += timedelta(weeks=1)

    max_journeys = max((len(w["bookings"]) for w in weeks), default=1)

    # ── Write Excel ───────────────────────────────────────────────────────────
    print(f"\n📊 Writing Excel file: {args.output}")
    write_excel(weeks, max_journeys, passenger_name, args.output)

    total_spend = sum(b["amount_raw"] for b in bookings)
    weeks_travelled = sum(1 for w in weeks if w["bookings"])
    print(f"\n{'─'*60}")
    print(f"📁 Saved: {args.output}")
    print(f"📅 Weeks analysed: {len(weeks)}")
    print(f"🚆 Weeks travelled: {weeks_travelled}")
    print(f"💶 Total spend ({passenger_name}'s share): €{total_spend:.2f}")
    print(f"{'─'*60}\n")


if __name__ == "__main__":
    main()
